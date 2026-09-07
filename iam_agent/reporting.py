"""Deterministic, authenticated report exports for conversational results."""

from __future__ import annotations

import asyncio
import csv
import html
import io
import json
import logging
import re
import time
from collections import OrderedDict
from dataclasses import dataclass
from datetime import UTC, date, datetime
from typing import Any
from uuid import uuid4

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    LongTable,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    TableStyle,
)

from iam_agent.models import ChatWorkflowResult, ReportDescriptor, ReportFormat

logger = logging.getLogger(__name__)

MIME_TYPES = {
    ReportFormat.CSV: "text/csv; charset=utf-8",
    ReportFormat.EXCEL: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ReportFormat.PDF: "application/pdf",
    ReportFormat.WORD: "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ReportFormat.POWERPOINT: (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    ),
}

FORMAT_PATTERNS = {
    ReportFormat.CSV: re.compile(r"\b(?:csv|comma[- ]separated)\b", re.IGNORECASE),
    ReportFormat.EXCEL: re.compile(r"\b(?:excel|exel|excell|xlsx?|spreadsheet)\b", re.IGNORECASE),
    ReportFormat.PDF: re.compile(r"\b(?:pdf|pfd)\b", re.IGNORECASE),
    ReportFormat.WORD: re.compile(
        r"\b(?:docx|word\s+(?:file|document|report|formats?)|(?:as|in|to)\s+word)\b",
        re.IGNORECASE,
    ),
    ReportFormat.POWERPOINT: re.compile(
        r"\b(?:pptx?|power\s?point|slides?|slide\s*deck|presentation)\b", re.IGNORECASE,
    ),
}


def requested_report_formats(text: str) -> list[ReportFormat]:
    """Extract explicit export formats without asking Gemini to control file generation."""
    if re.search(r"\b(?:all|every)\s+(?:report\s+)?formats?\b", text, re.IGNORECASE):
        return list(ReportFormat)
    matches = [
        report_format for report_format, pattern in FORMAT_PATTERNS.items() if pattern.search(text)
    ]
    # A bare "Word" is ambiguous in prose, but unambiguous beside another file format or an
    # export/report request. Accept the common "CSV, Excel, PDF and Word" phrasing.
    if ReportFormat.WORD not in matches and re.search(r"\b(?:word|workd)\b", text, re.IGNORECASE):
        has_export_context = bool(
            matches
            or re.search(r"\b(?:export|download|report|file|formats?)\b", text, re.IGNORECASE)
        )
        if has_export_context:
            matches.append(ReportFormat.WORD)
    return [report_format for report_format in ReportFormat if report_format in matches]


@dataclass(frozen=True)
class ReportArtifact:
    report_id: str
    owner: str
    format: ReportFormat
    filename: str
    mime_type: str
    content: bytes


@dataclass(frozen=True)
class _StoredReport:
    artifact: ReportArtifact
    expires_at: float


class ReportStore:
    """Small, bounded store for immediate downloads from the single Cloud Run instance."""

    def __init__(self, *, max_reports: int = 100, ttl_seconds: int = 15 * 60):
        self.max_reports = max_reports
        self.ttl_seconds = ttl_seconds
        self._reports: OrderedDict[str, _StoredReport] = OrderedDict()
        self._lock = asyncio.Lock()

    async def add(
        self,
        *,
        owner: str,
        report_format: ReportFormat,
        filename: str,
        content: bytes,
    ) -> ReportDescriptor:
        if len(content) > 16_000_000:
            raise ValueError("Report exceeds the 16 MB file limit; request fewer rows")
        report_id = uuid4().hex
        artifact = ReportArtifact(
            report_id=report_id,
            owner=owner.casefold(),
            format=report_format,
            filename=filename,
            mime_type=MIME_TYPES[report_format],
            content=content,
        )
        async with self._lock:
            self._remove_expired_locked()
            self._reports[report_id] = _StoredReport(
                artifact=artifact,
                expires_at=time.monotonic() + self.ttl_seconds,
            )
            while len(self._reports) > self.max_reports or sum(
                len(item.artifact.content) for item in self._reports.values()
            ) > 64_000_000:
                self._reports.popitem(last=False)
        return ReportDescriptor(
            report_id=report_id,
            format=report_format,
            filename=filename,
            mime_type=artifact.mime_type,
            download_url=f"/web/reports/{report_id}",
        )

    async def get(self, report_id: str, owner: str) -> ReportArtifact | None:
        async with self._lock:
            self._remove_expired_locked()
            stored = self._reports.get(report_id)
            if stored is None or stored.artifact.owner != owner.casefold():
                return None
            self._reports.move_to_end(report_id)
            return stored.artifact

    def _remove_expired_locked(self) -> None:
        now = time.monotonic()
        expired = [key for key, value in self._reports.items() if value.expires_at <= now]
        for key in expired:
            self._reports.pop(key, None)


class ReportExporter:
    def export(self, result: ChatWorkflowResult, report_format: ReportFormat) -> bytes:
        if report_format == ReportFormat.CSV:
            return self._csv(result)
        if report_format == ReportFormat.EXCEL:
            return self._excel(result)
        if report_format == ReportFormat.PDF:
            return self._pdf(result)
        if report_format == ReportFormat.WORD:
            return self._word(result)
        if report_format == ReportFormat.POWERPOINT:
            from iam_agent.slide_report import export_presentation

            return export_presentation(result)
        raise ValueError(f"Unsupported report format: {report_format}")

    def filename(self, result: ChatWorkflowResult, report_format: ReportFormat) -> str:
        objective = (
            result.plan.objective if result.plan else result.decomposition.clarified_question
        )
        slug = re.sub(r"[^a-z0-9]+", "-", objective.casefold()).strip("-")[:60]
        slug = slug or "iam-analysis"
        return f"{slug}-{datetime.now(UTC):%Y%m%d}.{report_format.value}"

    def _csv(self, result: ChatWorkflowResult) -> bytes:
        rows = self._flat_rows(result)
        fieldnames = self._fieldnames(rows)
        output = io.StringIO(newline="")
        writer = csv.DictWriter(output, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({key: self._spreadsheet_value(row.get(key, "")) for key in fieldnames})
        return output.getvalue().encode("utf-8-sig")

    def _excel(self, result: ChatWorkflowResult) -> bytes:
        workbook = Workbook()
        summary = workbook.active
        summary.title = "Summary"
        summary_rows = [
            ("Report generated", datetime.now(UTC).isoformat()),
            ("Question", result.decomposition.clarified_question),
            ("Answer", result.answer.answer),
            ("Domains", ", ".join(domain.value for domain in result.decomposition.domains)),
            ("Sources", ", ".join(result.answer.sources)),
        ]
        if result.report_document:
            summary_rows.extend([
                ("Report title", result.report_document.title),
                ("Executive summary", result.report_document.executive_summary),
                *((s.heading, s.body) for s in result.report_document.sections),
            ])
        for row in result.answer.key_findings:
            summary_rows.append(("Key finding", row))
        for row in result.answer.caveats:
            summary_rows.append(("Caveat", row))
        for row in summary_rows:
            summary.append([self._spreadsheet_value(value) for value in row])
        self._style_sheet(summary)
        summary.column_dimensions["A"].width = 26
        summary.column_dimensions["B"].width = 110
        for row in summary:
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            summary.row_dimensions[row[0].row].height = min(
                400, max(30, (len(str(row[1].value or "")) // 100 + 1) * 16),
            )

        scope = workbook.create_sheet("Query scope")
        scope.append(["Source", "Purpose", "Rows returned", "Filters", "Time window (days)"])
        if result.evidence:
            for item in result.evidence.results:
                scope.append([self._spreadsheet_value(value) for value in (
                    item.source_table, item.purpose, item.row_count,
                    [f.model_dump(mode="json") for f in item.applied_filters],
                    item.time_range_days if item.time_range_days is not None
                    else "All retained history",
                )])
        self._style_sheet(scope, header=True)

        if result.evidence:
            for index, query_result in enumerate(result.evidence.results, start=1):
                title = self._sheet_title(query_result.concept, index, workbook.sheetnames)
                sheet = workbook.create_sheet(title)
                rows = query_result.rows or [{"message": "No rows returned"}]
                fields = self._fieldnames(rows)
                sheet.append(fields)
                for row in rows:
                    sheet.append([self._spreadsheet_value(row.get(field, "")) for field in fields])
                self._style_sheet(sheet, header=True)

        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    def _word(self, result: ChatWorkflowResult) -> bytes:
        document = Document()
        document.add_heading(
            result.report_document.title if result.report_document else "IAM Analytics Report", 0,
        )
        document.add_paragraph(f"Generated: {datetime.now(UTC).isoformat()}")
        document.add_heading("Question", level=1)
        document.add_paragraph(result.decomposition.clarified_question)
        document.add_heading("Analysis", level=1)
        document.add_paragraph(result.answer.answer)
        if result.report_document:
            document.add_heading("Executive summary", level=1)
            document.add_paragraph(result.report_document.executive_summary)
            for section in result.report_document.sections:
                document.add_heading(section.heading, level=1)
                document.add_paragraph(section.body)
        self._word_list(document, "Key findings", result.answer.key_findings)
        self._word_list(document, "Sources", result.answer.sources)
        self._word_list(document, "Caveats", result.answer.caveats)

        if result.evidence:
            document.add_heading("Evidence", level=1)
            for query_result in result.evidence.results:
                document.add_heading(query_result.purpose, level=2)
                document.add_paragraph(
                    f"Concept: {query_result.concept} | Source: {query_result.source_table} | "
                    f"Rows: {query_result.row_count}"
                )
                if not query_result.rows:
                    document.add_paragraph("No rows returned.")
                    continue
                fields = self._fieldnames(query_result.rows)
                for start in range(0, len(fields), 4):
                    group = fields[start:start + 4]
                    table = document.add_table(rows=1, cols=len(group) + 1)
                    table.style = "Light Shading Accent 1"
                    for cell, field in zip(table.rows[0].cells, ["Row", *group], strict=True):
                        cell.text = field
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
                    for index, row in enumerate(query_result.rows, 1):
                        cells = table.add_row().cells
                        cells[0].text = str(index)
                        for cell, field in zip(cells[1:], group, strict=True):
                            cell.text = self._display_value(row.get(field, ""))
                    document.add_paragraph()

        output = io.BytesIO()
        document.save(output)
        return output.getvalue()

    def _pdf(self, result: ChatWorkflowResult) -> bytes:
        output = io.BytesIO()
        styles = getSampleStyleSheet()
        document = SimpleDocTemplate(
            output,
            pagesize=landscape(letter),
            rightMargin=0.35 * inch,
            leftMargin=0.35 * inch,
            topMargin=0.4 * inch,
            bottomMargin=0.4 * inch,
            title="IAM Analytics Report",
        )
        story: list[Any] = [
            Paragraph(html.escape(
                result.report_document.title if result.report_document else "IAM Analytics Report"
            ), styles["Title"]),
            Paragraph(f"Generated: {datetime.now(UTC).isoformat()}", styles["Normal"]),
            Spacer(1, 8),
            Paragraph("Question", styles["Heading2"]),
            Paragraph(html.escape(result.decomposition.clarified_question), styles["BodyText"]),
            Paragraph("Analysis", styles["Heading2"]),
            Paragraph(html.escape(result.answer.answer), styles["BodyText"]),
        ]
        if result.report_document:
            story.extend([
                Paragraph("Executive summary", styles["Heading2"]),
                Paragraph(
                    html.escape(result.report_document.executive_summary), styles["BodyText"],
                ),
            ])
            for section in result.report_document.sections:
                story.append(Paragraph(html.escape(section.heading), styles["Heading2"]))
                for paragraph in section.body.splitlines():
                    if paragraph.strip():
                        story.append(Paragraph(html.escape(paragraph), styles["BodyText"]))
        for heading, values in (
            ("Key findings", result.answer.key_findings),
            ("Sources", result.answer.sources),
            ("Caveats", result.answer.caveats),
        ):
            if values:
                story.append(Paragraph(heading, styles["Heading2"]))
                story.extend(
                    Paragraph(f"• {html.escape(self._display_value(value))}", styles["BodyText"])
                    for value in values
                )

        if result.evidence:
            for query_result in result.evidence.results:
                story.extend(
                    [
                        PageBreak(),
                        Paragraph(html.escape(query_result.purpose), styles["Heading2"]),
                        Paragraph(
                            html.escape(
                                f"Concept: {query_result.concept} | "
                                f"Source: {query_result.source_table} | "
                                f"Rows: {query_result.row_count}"
                            ),
                            styles["BodyText"],
                        ),
                        Spacer(1, 6),
                    ]
                )
                if not query_result.rows:
                    story.append(Paragraph("No rows returned.", styles["BodyText"]))
                    continue
                fields = self._fieldnames(query_result.rows)
                story.append(Paragraph(
                    "All returned rows are included. Long cells are abbreviated; use Excel or CSV "
                    "for complete cell values. Row numbers align column groups.",
                    styles["BodyText"],
                ))
                for start in range(0, len(fields), 5):
                    group = fields[start:start + 5]
                    table_data = [["Row", *[
                        self._pdf_cell(field, styles["BodyText"]) for field in group
                    ]]]
                    table_data.extend(
                        [str(index), *[
                            self._pdf_cell(row.get(field, ""), styles["BodyText"])
                            for field in group
                        ]] for index, row in enumerate(query_result.rows, 1)
                    )
                    page_width = landscape(letter)[0] - 0.7 * inch
                    table = LongTable(
                        table_data,
                        colWidths=[25] + [(page_width - 25) / len(group)] * len(group),
                        repeatRows=1,
                    )
                    table.setStyle(self._pdf_table_style())
                    story.extend([Spacer(1, 8), table])
        document.build(story)
        return output.getvalue()

    def _flat_rows(self, result: ChatWorkflowResult) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = [
            {
                "_record_type": "question",
                "text": result.decomposition.clarified_question,
            },
            {"_record_type": "summary", "text": result.answer.answer},
            *({"_record_type": "finding", "text": value} for value in result.answer.key_findings),
            *({"_record_type": "source", "text": value} for value in result.answer.sources),
            *({"_record_type": "caveat", "text": value} for value in result.answer.caveats),
        ]
        if result.report_document:
            rows.extend([
                {
                    "_record_type": "executive_summary",
                    "text": result.report_document.executive_summary,
                },
                *({"_record_type": "report_section", "_purpose": s.heading, "text": s.body}
                  for s in result.report_document.sections),
            ])
        if result.evidence:
            for query_result in result.evidence.results:
                for evidence_row in query_result.rows:
                    rows.append(
                        {
                            "_record_type": "evidence",
                            "_step_id": query_result.step_id,
                            "_purpose": query_result.purpose,
                            "_concept": query_result.concept,
                            "_source_table": query_result.source_table,
                            **evidence_row,
                        }
                    )
        return rows

    @staticmethod
    def _fieldnames(rows: list[dict[str, Any]]) -> list[str]:
        preferred = ["_record_type", "_step_id", "_purpose", "_concept", "_source_table", "text"]
        available = {key for row in rows for key in row}
        return [key for key in preferred if key in available] + sorted(available - set(preferred))

    @staticmethod
    def _display_value(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, (date, datetime)):
            return value.isoformat()
        if isinstance(value, (dict, list, tuple)):
            return json.dumps(value, default=str, ensure_ascii=False, sort_keys=True)
        return str(value)

    @classmethod
    def _spreadsheet_value(cls, value: Any) -> str | int | float | bool:
        if isinstance(value, bool) or isinstance(value, (int, float)):
            return value
        rendered = cls._display_value(value)
        if rendered.startswith(("=", "+", "-", "@")):
            return "'" + rendered
        return rendered

    @classmethod
    def _pdf_cell(cls, value: Any, style) -> Paragraph:
        rendered = cls._display_value(value)
        if len(rendered) > 500:
            rendered = rendered[:497] + "..."
        return Paragraph(html.escape(rendered), style)

    @staticmethod
    def _style_sheet(sheet, *, header: bool = False) -> None:
        sheet.freeze_panes = "A2" if header else None
        if header:
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2758D7")
            sheet.auto_filter.ref = sheet.dimensions
        else:
            for cell in sheet[1]:
                cell.font = Font(bold=True)
        for column in sheet.columns:
            width = min(60, max(12, *(len(str(cell.value or "")) + 2 for cell in column)))
            sheet.column_dimensions[get_column_letter(column[0].column)].width = width

    @staticmethod
    def _sheet_title(concept: str, index: int, existing: list[str]) -> str:
        base = re.sub(r"[\\/*?:\[\]]", "_", concept)[:25] or f"Evidence {index}"
        candidate = base
        suffix = 2
        while candidate in existing:
            candidate = f"{base[:27]} {suffix}"
            suffix += 1
        return candidate[:31]

    @staticmethod
    def _word_list(document: Document, heading: str, values: list[str]) -> None:
        if not values:
            return
        document.add_heading(heading, level=1)
        for value in values:
            document.add_paragraph(value, style="List Bullet")

    @staticmethod
    def _pdf_table_style() -> TableStyle:
        return TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2758D7")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#DCE3ED")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F6FA")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ]
        )


class ReportService:
    def __init__(self, store: ReportStore, exporter: ReportExporter | None = None):
        self.store = store
        self.exporter = exporter or ReportExporter()
        self._render_lock = asyncio.Lock()

    async def create_reports(
        self,
        *,
        owner: str,
        result: ChatWorkflowResult,
        formats: list[ReportFormat],
    ) -> tuple[list[ReportDescriptor], list[str]]:
        descriptors: list[ReportDescriptor] = []
        failures: list[str] = []
        for report_format in dict.fromkeys(formats):
            try:
                async with self._render_lock:
                    content = await asyncio.to_thread(self.exporter.export, result, report_format)
                filename = self.exporter.filename(result, report_format)
                descriptors.append(
                    await self.store.add(
                        owner=owner,
                        report_format=report_format,
                        filename=filename,
                        content=content,
                    )
                )
            except Exception:
                logger.exception("Failed to generate %s report", report_format.value)
                failures.append(report_format.value)
        return descriptors, failures
