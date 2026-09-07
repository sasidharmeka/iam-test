"""Editable PowerPoint reports built from the same saved analytical evidence."""

from __future__ import annotations

import io
import textwrap
from datetime import UTC, datetime

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

from iam_agent.models import ChatWorkflowResult

INK = RGBColor.from_string("162238")
BLUE = RGBColor.from_string("2758D7")
MUTED = RGBColor.from_string("65738A")


def export_presentation(result: ChatWorkflowResult) -> bytes:
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)

    def textbox(slide, text, x, y, w, h, size, color=INK, bold=False):
        frame = slide.shapes.add_textbox(
            Inches(x), Inches(y), Inches(w), Inches(h),
        ).text_frame
        frame.word_wrap = True
        for i, line in enumerate(text.split("\n")):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.name = "Aptos"
            paragraph.font.size, paragraph.font.color.rgb = Pt(size), color
            paragraph.font.bold = bold
            paragraph.space_after = Pt(8)

    def slide(title):
        page = deck.slides.add_slide(deck.slide_layouts[6])
        textbox(page, "IAM ANALYTICS  /  EVIDENCE BRIEF", .6, .28, 12, .3, 11, BLUE, True)
        textbox(page, title, .6, .75, 12, 1.2, 28 if len(title) < 70 else 24, INK, True)
        textbox(
            page, f"Generated {datetime.now(UTC):%d %b %Y}  ·  {len(deck.slides)}",
            .6, 7, 12, .3, 10, MUTED,
        )
        return page

    def text_pages(title, text):
        # Explicit pagination prevents long model paragraphs from overflowing slides.
        lines = []
        for paragraph in text.splitlines():
            lines.extend(textwrap.wrap(paragraph, width=78) or [""])
        for start in range(0, max(1, len(lines)), 9):
            page = slide(title if start == 0 else title + " (continued)")
            textbox(page, "\n".join(lines[start:start + 9]), .65, 2.1, 12, 4.6, 20)

    document = result.report_document
    text_pages(
        document.title if document else "IAM Analytics Report",
        result.decomposition.clarified_question,
    )
    text_pages(
        "Executive summary", document.executive_summary if document else result.answer.answer,
    )
    if result.answer.key_findings:
        text_pages("Key findings", "\n".join(result.answer.key_findings))
    if result.evidence:
        for evidence in result.evidence.results:
            rows = evidence.rows
            metric = next((
                key for key in ("risk_score", "maximum_risk_score")
                if rows and all(isinstance(r.get(key), (int, float)) for r in rows)
            ), None)
            dimension = next((
                key for key in ("user", "user_id") if rows and all(key in r for r in rows)
            ), None)
            if metric and dimension:
                top = sorted(rows, key=lambda r: r[metric], reverse=True)[:10]
                data = CategoryChartData()
                data.categories = [str(r[dimension]) for r in top]
                data.add_series(metric.replace("_", " "), [r[metric] for r in top])
                page = slide("Highest scores in the returned evidence")
                page.shapes.add_chart(
                    XL_CHART_TYPE.BAR_CLUSTERED, Inches(.6), Inches(1.9),
                    Inches(12), Inches(4.5), data,
                )
                textbox(
                    page, f"Source: {evidence.source_table} · Top {len(top)} returned rows; "
                    "historical scores do not establish current risk.",
                    .65, 6.5, 12, .4, 11, MUTED,
                )
                break
    if document:
        for section in document.sections:
            text_pages(section.heading, section.body)
    scope = []
    if result.evidence:
        for evidence in result.evidence.results:
            filters = "; ".join(
                f"{f.field} {f.operator.value} {f.value}" for f in evidence.applied_filters
            ) or "No explicit field filters"
            window = (
                f"{evidence.time_range_days} days" if evidence.time_range_days
                else "All retained history"
            )
            scope.append(
                f"{evidence.source_table}: {evidence.row_count} returned rows. {window}. {filters}"
            )
    text_pages(
        "Sources and scope", "\n".join(scope or result.answer.sources)
        or "General IAM guidance; no project evidence queried.",
    )
    text_pages("Limitations", "\n".join(result.answer.caveats + [
        "This deck summarizes the saved analysis. Request Excel or CSV for all returned rows. "
        "An export does not perform a new query or retrieve rows outside the original limits."
    ]))
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()
