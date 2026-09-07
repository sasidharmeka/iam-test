from __future__ import annotations

import io
import json

import pytest
from openpyxl import load_workbook

from iam_agent.actions import ActionResolver, ReportAuthor
from iam_agent.conversation import ConversationStore
from iam_agent.conversation_actions import ConversationActions
from iam_agent.email_service import EmailService
from iam_agent.models import (
    ConversationAction,
    EmailContent,
    ReportDocument,
    ReportFormat,
    ReportSection,
)
from iam_agent.reporting import ReportService, ReportStore
from iam_agent.result_store import ResultStore
from test_reporting import sample_result

OWNER, KEY = "analyst@example.com", "one"


class ActionModel:
    def __init__(self):
        self.prompts = []
        self.decision = ConversationAction(action="analyze")

    async def generate_structured(self, *, prompt, response_model, **kwargs):
        self.prompts.append(json.loads(prompt))
        if response_model is ConversationAction:
            return self.decision.model_copy(deep=True)
        if response_model is ReportDocument:
            return ReportDocument(
                title="Identity risk review",
                executive_summary="user12 has an observed score of 91.",
                sections=[ReportSection(heading="Scope", body="The saved historical evidence.")],
            )
        if response_model is EmailContent:
            return EmailContent(
                subject="Identity risk review", body="Please review the risk analysis.",
            )
        raise AssertionError(response_model)


def build_actions():
    model, files, results = ActionModel(), ReportStore(), ResultStore()
    actions = ConversationActions(
        resolver=ActionResolver(model), author=ReportAuthor(model), results=results,
        reports=ReportService(files), emails=EmailService(files),
        conversations=ConversationStore(max_sessions=10, max_turns=6),
    )
    return actions, model


async def ask(actions, question, source=""):
    return await actions.prepare(
        owner=OWNER, key=KEY, question=question, history=[], formats=[], source_result_id=source,
    )


@pytest.mark.asyncio
@pytest.mark.parametrize("question", [
    "ok give me this data so i can download it",
    "can i downlaod this data", "Export it as Excel", "PDF please", "prepare this as a ppt",
])
async def test_export_uses_exact_previous_evidence_without_domain_routing(question):
    actions, model = build_actions()
    result = sample_result()
    await actions.results.save(OWNER, KEY, result)
    _, _, response = await ask(actions, question)
    assert response.result_id == result.result_id
    assert response.reports
    assert response.evidence.results[0].rows == result.evidence.results[0].rows
    assert not any("available_results" in p for p in model.prompts)
    saved = await actions.results.get(result.result_id, OWNER, KEY)
    assert saved.answer.answer == result.answer.answer
    if response.reports[0].format == ReportFormat.EXCEL:
        artifact = await actions.reports.store.get(response.reports[0].report_id, OWNER)
        workbook = load_workbook(io.BytesIO(artifact.content))
        assert 91 in [cell.value for row in workbook["risk_scores"] for cell in row]


@pytest.mark.asyncio
async def test_new_topic_with_format_is_not_mistaken_for_export():
    actions, model = build_actions()
    result = sample_result()
    await actions.results.save(OWNER, KEY, result)
    model.decision = ConversationAction(
        action="analyze", resolved_question="Show MFA bypasses as PDF", formats=[ReportFormat.PDF],
    )
    question, formats, response = await ask(actions, "Show MFA as PDF")
    assert response is None
    assert question == "Show MFA bypasses as PDF"
    assert formats == [ReportFormat.PDF]
    assert model.prompts[0]["available_results"][0]["plan"]


@pytest.mark.asyncio
async def test_previous_excel_attachment_resolves_older_result_not_newest_analysis():
    actions, model = build_actions()
    first = sample_result()
    await actions.results.save(OWNER, KEY, first)
    _, _, exported = await ask(actions, "download this as Excel")
    expected_id = exported.reports[0].report_id
    second = sample_result()
    second.answer.answer = "An unrelated MFA analysis."
    await actions.results.save(OWNER, KEY, second)
    _, _, response = await ask(
        actions, "Write a Gmail and attach the previous Excel you produced",
    )
    assert response.email_draft["recipient"] == ""
    assert response.email_draft["state"] == "draft"
    assert response.email_draft["attachment_ids"] == [expected_id]
    assert response.result_id == first.result_id
    assert model.prompts[-1]["analysis"]["answer"] == first.answer.answer


@pytest.mark.asyncio
async def test_result_reference_checks_owner_conversation_and_expiry():
    actions, model = build_actions()
    result = sample_result()
    await actions.results.save(OWNER, KEY, result)
    assert await actions.results.get(result.result_id, "other@example.com", KEY) is None
    assert await actions.results.get(result.result_id, OWNER, "different") is None
    actions.results._items[result.result_id].expires_at = 0
    _, _, response = await ask(actions, "download this", result.result_id)
    assert "expired" in response.answer.answer
    assert not response.reports


@pytest.mark.asyncio
async def test_clicking_older_suggestion_uses_its_result_anchor():
    actions, model = build_actions()
    first, second = sample_result(), sample_result()
    await actions.results.save(OWNER, KEY, first)
    second.evidence.results[0].rows[0]["maximum_risk_score"] = 10
    await actions.results.save(OWNER, KEY, second)
    _, _, response = await ask(actions, "download this", first.result_id)
    assert response.result_id == first.result_id
    assert response.evidence.results[0].rows[0]["maximum_risk_score"] == 91


@pytest.mark.asyncio
async def test_email_recipient_must_be_explicit_not_model_invented():
    actions, model = build_actions()
    model.decision = ConversationAction(action="email", recipient="invented@example.com")
    result = sample_result()
    await actions.results.save(OWNER, KEY, result)
    _, _, response = await ask(actions, "Please compose an email about this")
    assert response.email_draft["recipient"] == ""
    _, _, response = await ask(actions, "Write an email to requested@example.com about this")
    assert response.email_draft["recipient"] == "requested@example.com"


@pytest.mark.asyncio
async def test_unknown_attachment_id_is_rejected():
    actions, model = build_actions()
    result = sample_result()
    await actions.results.save(OWNER, KEY, result)
    model.decision = ConversationAction(action="email", attachment_ids=["invented-file"])
    _, _, response = await ask(actions, "Can you share this with the team?")
    assert response.email_draft is None
    assert "attachment" in response.answer.answer.lower()


@pytest.mark.asyncio
async def test_full_orchestrator_analytics_then_export_then_email_queries_once(tmp_path):
    from iam_agent.models import ChatActor
    from test_orchestrator import FakeMCPClient, FakeModel, build_orchestrator

    actions, model = build_actions()
    mcp = FakeMCPClient()
    orchestrator = build_orchestrator(tmp_path, FakeModel(data_question=True), mcp)
    orchestrator.actions = actions
    actor = ChatActor(principal=OWNER)
    initial = await orchestrator.answer(
        actor=actor, question="Which users are riskiest?", conversation_key=KEY, allow_actions=True,
    )
    exported = await orchestrator.answer(
        actor=actor, question="give this data so i can download it",
        conversation_key=KEY, allow_actions=True,
    )
    draft = await orchestrator.answer(
        actor=actor, question="email this with the previous Excel",
        conversation_key=KEY, allow_actions=True,
    )
    assert len(mcp.plans) == 1
    assert initial.result_id == exported.result_id == draft.result_id
    assert draft.email_draft["attachment_ids"] == [exported.reports[0].report_id]
