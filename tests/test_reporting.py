from __future__ import annotations

import io

import pytest
from docx import Document
from openpyxl import load_workbook

from iam_agent.models import (
    AnalysisPlan,
    AnalystAnswer,
    AnswerMode,
    ChatWorkflowResult,
    EvidenceBundle,
    IAMDomain,
    QueryOperation,
    QueryStep,
    QueryStepResult,
    QuestionDecomposition,
    ReportFormat,
)
from iam_agent.reporting import ReportExporter, ReportStore, requested_report_formats


def sample_result() -> ChatWorkflowResult:
    plan = AnalysisPlan(
        objective="Rank risky users",
        steps=[
            QueryStep(
                step_id="rank_risk",
                purpose="Rank observed user risk",
                concept="risk_scores",
                operation=QueryOperation.RANK,
                metrics=["maximum_risk_score"],
                dimensions=["user"],
            )
        ],
    )
    return ChatWorkflowResult(
        decomposition=QuestionDecomposition(
            answer_mode=AnswerMode.DATA,
            clarified_question="Who are the riskiest users?",
            domains=[IAMDomain.IDENTITY_RISK],
            requires_project_data=True,
        ),
        plan=plan,
        evidence=EvidenceBundle(
            objective=plan.objective,
            queried_tables=["risk_scores"],
            results=[
                QueryStepResult(
                    step_id="rank_risk",
                    purpose="Rank observed user risk",
                    concept="risk_scores",
                    source_table="risk_scores",
                    rows=[
                        {
                            "user": "user12",
                            "maximum_risk_score": 91,
                            "note": '=HYPERLINK("https://example.invalid")',
                        }
                    ],
                    row_count=1,
                )
            ],
        ),
        answer=AnalystAnswer(
            answer="user12 has the highest observed score.",
            key_findings=["user12: 91"],
            sources=["risk_scores"],
            caveats=["This reflects the available output."],
        ),
    )


def test_requested_report_formats_are_deterministic_and_deduplicated() -> None:
    assert requested_report_formats("Export this as CSV, Excel, PDF, and a Word document") == [
        ReportFormat.CSV,
        ReportFormat.EXCEL,
        ReportFormat.PDF,
        ReportFormat.WORD,
    ]
    assert requested_report_formats("Explain the word risk") == []
    assert requested_report_formats("Export as csv, exel, pfd and workd") == [
        ReportFormat.CSV, ReportFormat.EXCEL, ReportFormat.PDF, ReportFormat.WORD,
    ]
    assert requested_report_formats("Prepare a PowerPoint") == [ReportFormat.POWERPOINT]
    assert requested_report_formats("Give me all report formats") == list(ReportFormat)


def test_all_report_formats_are_valid_files_and_excel_cells_are_safe() -> None:
    exporter = ReportExporter()
    result = sample_result()

    csv_content = exporter.export(result, ReportFormat.CSV)
    assert csv_content.startswith(b"\xef\xbb\xbf")
    assert b"'=HYPERLINK" in csv_content

    excel_content = exporter.export(result, ReportFormat.EXCEL)
    workbook = load_workbook(io.BytesIO(excel_content))
    assert workbook.sheetnames == ["Summary", "Query scope", "risk_scores"]
    evidence_values = [cell.value for row in workbook["risk_scores"].iter_rows() for cell in row]
    assert '\'=HYPERLINK("https://example.invalid")' in evidence_values

    pdf_content = exporter.export(result, ReportFormat.PDF)
    assert pdf_content.startswith(b"%PDF")

    word_content = exporter.export(result, ReportFormat.WORD)
    document = Document(io.BytesIO(word_content))
    assert document.paragraphs[0].text == "IAM Analytics Report"

    from pptx import Presentation

    deck = Presentation(io.BytesIO(exporter.export(result, ReportFormat.POWERPOINT)))
    assert len(deck.slides) >= 5
    chart = next(
        shape.chart for slide in deck.slides for shape in slide.shapes if shape.has_chart
    )
    assert chart.series[0].values == (91.0,)
    assert "user12" in [c.label for c in chart.plots[0].categories]


@pytest.mark.asyncio
async def test_report_store_does_not_leak_files_between_users() -> None:
    store = ReportStore()
    descriptor = await store.add(
        owner="analyst@example.com",
        report_format=ReportFormat.PDF,
        filename="report.pdf",
        content=b"%PDF-demo",
    )

    assert await store.get(descriptor.report_id, "other@example.com") is None
    artifact = await store.get(descriptor.report_id, "ANALYST@example.com")
    assert artifact is not None
    assert artifact.content == b"%PDF-demo"
